"""Build a report from reviewed JSON; equations are native editable Word OMML.

No extraction, scientific inference, or numerical calculation is performed here.
"""
import argparse
import copy
import json
import hashlib
import math
import re
import unicodedata
from itertools import groupby
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
TEMPLATE = Path(__file__).resolve().parents[1] / 'assets' / 'physics-report-template.docx'
HEADINGS = ['一、实验目的', '二、实验原理', '三、实验仪器', '四、实验步骤',
            '五、实验数据记录', '六、数据处理（绘图及计算）', '七、实验结果分析',
            '八、课后思考题及实验拓展']


def set_run_fonts(properties):
    fonts = properties.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        properties.insert(0, fonts)
    fonts.attrib.clear()
    for key in ('ascii', 'hAnsi', 'cs'):
        fonts.set(qn('w:' + key), 'Times New Roman')
    fonts.set(qn('w:eastAsia'), '宋体')



def punctuation_groups(text):
    return [(punct, ''.join(chars)) for punct, chars in
            groupby(text, lambda c: unicodedata.category(c).startswith('P'))]


def apply_punctuation_fonts(properties, punctuation):
    set_run_fonts(properties)
    if punctuation:
        fonts = properties.find(qn('w:rFonts'))
        for key in ('ascii', 'hAnsi', 'cs', 'eastAsia'):
            fonts.set(qn('w:' + key), '宋体')


def normalize_math_fonts(expr, inherited=None):
    for run in list(expr.iter(qn('m:r'))):
        text = run.find(qn('m:t'))
        if text is None:
            continue
        parent = run.getparent()
        offset = parent.index(run)
        for punctuation, value in punctuation_groups(text.text or ''):
            child = copy.deepcopy(run)
            child.find(qn('m:t')).text = value
            math_properties = child.find(qn('m:rPr'))
            if math_properties is None:
                math_properties = element('rPr')
                child.insert(0, math_properties)
            if math_properties.find(qn('m:nor')) is None:
                math_properties.append(element('nor'))
            properties = child.find(qn('w:rPr'))
            if properties is None:
                properties = copy.deepcopy(inherited) if inherited is not None else OxmlElement('w:rPr')
                child.insert(1, properties)
            apply_punctuation_fonts(properties, punctuation)
            parent.insert(offset, child)
            offset += 1
        parent.remove(run)


def inline_segments(text):
    # Typed inline expressions prevent guessing whether Udy means U with subscript dy.
    decoder = json.JSONDecoder()
    letters = re.compile(r'[A-Za-z\u00c0-\u024f\u0370-\u03ff][A-Za-z0-9_\u00c0-\u024f\u0370-\u03ff]*')
    cursor = 0
    while cursor < len(text):
        marker = text.find('{{math:', cursor)
        end = marker if marker >= 0 else len(text)
        plain = text[cursor:end]
        last = 0
        for match in letters.finditer(plain):
            if match.start() > last:
                yield False, plain[last:match.start()]
            yield True, match.group()
            last = match.end()
        if last < len(plain):
            yield False, plain[last:]
        if marker < 0:
            break
        start = marker + len('{{math:')
        start += len(text[start:]) - len(text[start:].lstrip())
        value, consumed = decoder.raw_decode(text[start:])
        cursor = start + consumed
        if text[cursor:cursor+2] != '}}':
            raise ValueError('Inline math requires a JSON expression followed by }}')
        cursor += 2
        yield True, value


def format_document_text(doc):
    for run in list(doc.element.iter(qn('w:r'))):
        text = run.find(qn('w:t'))
        if text is None:
            continue
        # Generated text runs must not also contain drawings or other run content.
        if any(c.tag not in {qn('w:rPr'), qn('w:t')} for c in run):
            continue
        parent, offset = run.getparent(), run.getparent().index(run)
        properties = run.find(qn('w:rPr'))
        for is_math, value in inline_segments(text.text or ''):
            if is_math:
                expr = math_node(value)
                if expr.tag != qn('m:oMath'):
                    if isinstance(value, dict) and 'seq' in value:
                        expr = element('oMath', list(expr))
                    else:
                        expr = element('oMath', [expr])
                normalize_math_fonts(expr, properties)
                parent.insert(offset, expr)
                offset += 1
            else:
                for punctuation, piece in punctuation_groups(value):
                    child = copy.deepcopy(run)
                    child.find(qn('w:t')).text = piece
                    child.find(qn('w:t')).set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                    props = child.find(qn('w:rPr'))
                    if props is None:
                        props = OxmlElement('w:rPr')
                        child.insert(0, props)
                    apply_punctuation_fonts(props, punctuation)
                    parent.insert(offset, child)
                    offset += 1
        parent.remove(run)
    # Display equations also need Songti punctuation, including ASCII punctuation.
    for expr in list(doc.element.iter(qn('m:oMath'))):
        normalize_math_fonts(expr)


def configure_styles(doc):
    """User-specified typography, independent of inherited template defaults."""
    specs = {'Normal': (12, False, False, WD_ALIGN_PARAGRAPH.JUSTIFY, True),
             'Title': (22, True, False, WD_ALIGN_PARAGRAPH.CENTER, False),
             'PLR Info': (12, False, False, WD_ALIGN_PARAGRAPH.CENTER, False),
             'PLR Figure': (12, False, False, WD_ALIGN_PARAGRAPH.CENTER, False),
             'PLR Caption': (10.5, True, False, WD_ALIGN_PARAGRAPH.CENTER, False),
             'PLR Note': (9, True, False, WD_ALIGN_PARAGRAPH.JUSTIFY, False),
             'PLR Table': (10.5, False, False, WD_ALIGN_PARAGRAPH.CENTER, False),
             'PLR Emphasis': (12, True, False, WD_ALIGN_PARAGRAPH.JUSTIFY, True),
             'PLR Equation': (12, False, False, WD_ALIGN_PARAGRAPH.CENTER, False)}
    for level, size, bold, italic in [(1, 15, True, False), (2, 14, True, False),
                                      (3, 12, True, False), (4, 12, False, True), (5, 12, True, True)]:
        specs['Heading ' + str(level)] = (size, bold, italic, WD_ALIGN_PARAGRAPH.LEFT, False)
    for name, (size, bold, italic, alignment, indent) in specs.items():
        style = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        if name != 'Normal':
            style.base_style = doc.styles['Normal']
        style.font.name = 'Times New Roman'
        style.font.size, style.font.bold, style.font.italic = Pt(size), bold, italic
        style.font.color.rgb = RGBColor(0, 0, 0)
        set_run_fonts(style.element.get_or_add_rPr())
        fmt = style.paragraph_format
        fmt.alignment = alignment
        fmt.space_before, fmt.space_after = Pt(0), Pt(4)
        fmt.line_spacing = 1.25
        fmt.keep_with_next = name.startswith('Heading ')
        fmt.keep_together = False
        fmt.page_break_before = False
        fmt.widow_control = True
        ppr = style.element.get_or_add_pPr()
        old = ppr.find(qn('w:ind'))
        if old is not None:
            ppr.remove(old)
        ind = OxmlElement('w:ind')
        ind.set(qn('w:left'), '0')
        ind.set(qn('w:right'), '0')
        ind.set(qn('w:firstLineChars'), '200' if indent else '0')
        ind.set(qn('w:firstLine'), '0')
        ppr.append(ind)
        if name.startswith('Heading '):
            fmt.space_before, fmt.space_after = Pt(8), Pt(4)
        if name in {'PLR Caption', 'PLR Note', 'PLR Table'}:
            fmt.line_spacing = 1.0
            fmt.space_after = Pt(3)
    # Explicit fonts also cover template styles such as headers and footers.
    for style in doc.styles:
        set_run_fonts(style.element.get_or_add_rPr())


def chinese_label(value, field):
    if value is not None and (not isinstance(value, str) or
                              (value.strip() and not re.search('[\u3400-\u9fff]', value))):
        raise ValueError(field + ' must use Chinese text; symbols and units may remain Western')


def three_line_table(table, has_note=False):
    table.style = 'Normal Table'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        line = OxmlElement('w:' + edge)
        line.set(qn('w:val'), 'single' if edge in {'top', 'bottom'} else 'nil')
        line.set(qn('w:sz'), '12')
        line.set(qn('w:color'), '000000')
        borders.append(line)
    old = table._tbl.tblPr.find(qn('w:tblBorders'))
    if old is not None:
        table._tbl.tblPr.remove(old)
    table._tbl.tblPr.append(borders)
    for index, row in enumerate(table.rows):
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            old = tcpr.find(qn('w:tcBorders'))
            if old is not None:
                tcpr.remove(old)
            edges = OxmlElement('w:tcBorders')
            for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                line = OxmlElement('w:' + edge)
                line.set(qn('w:val'), 'single' if index == 0 and edge == 'bottom' else 'nil')
                line.set(qn('w:sz'), '6')
                line.set(qn('w:color'), '000000')
                edges.append(line)
            # Table outer borders supply top/bottom; avoid overriding them with nil cell edges.
            for edge in ('top', 'bottom'):
                if (index == 0 and edge == 'top') or (index == len(table.rows)-1 and edge == 'bottom'):
                    node = edges.find(qn('w:' + edge))
                    edges.remove(node)
            tcpr.append(edges)
            for p in cell.paragraphs:
                p.style = 'PLR Table'
                p.paragraph_format.keep_with_next = index == 0 or (has_note and index == len(table.rows)-1)
                for run in p.runs:
                    run.bold = index == 0


def element(tag, children=()):
    node = OxmlElement('m:' + tag)
    for child in children:
        node.append(child)
    return node


def math_node(value):
    """Convert a small typed expression tree to OMML; fail on unknown syntax."""
    if isinstance(value, (str, int, float)):
        if isinstance(value, bool) or (isinstance(value, str) and not value.strip()) or (isinstance(value, (int, float)) and not math.isfinite(value)):
            raise ValueError('Empty, boolean or non-finite equation value')
        text = element('t')
        text.text = str(value)
        text.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        return element('r', [text])
    if not isinstance(value, dict) or len(value) != 1:
        raise ValueError('Equation must be a scalar or one-key expression object')
    key, args = next(iter(value.items()))
    if key == 'omml':
        if not isinstance(args, str) or '<!DOCTYPE' in args.upper() or '<!ENTITY' in args.upper():
            raise ValueError('OMML must be XML without a DTD or entities')
        try:
            node = etree.fromstring(args.encode('utf-8'), etree.XMLParser(resolve_entities=False, no_network=True))
        except etree.XMLSyntaxError as exc:
            raise ValueError('Invalid OMML XML') from exc
        if node.tag != qn('m:oMath') or not any((t.text or '').strip() for t in node.findall('.//m:t', NS)):
            raise ValueError('OMML requires a non-empty m:oMath root')
        allowed = {NS['m'], NS['w']}
        if any(not isinstance(n.tag, str) or etree.QName(n).namespace not in allowed for n in node.iter()):
            raise ValueError('OMML contains unsupported XML elements')
        return node
    if key == 'seq' and isinstance(args, list) and args:
        return element('e', [child for x in args for child in math_children(x)])
    specs = {'frac': ('f', ['num', 'den']), 'sub': ('sSub', ['e', 'sub']),
             'sup': ('sSup', ['e', 'sup']), 'subsup': ('sSubSup', ['e', 'sub', 'sup'])}
    if key in specs:
        tag, slots = specs[key]
        if not isinstance(args, list) or len(args) != len(slots):
            raise ValueError('Wrong arity for equation operation ' + key)
        return element(tag, [element(slot, math_children(arg)) for slot, arg in zip(slots, args)])
    if key == 'sqrt':
        prop = element('radPr')
        hidden = element('degHide')
        hidden.set(qn('m:val'), '1')
        prop.append(hidden)
        return element('rad', [prop, element('deg'), element('e', math_children(args))])
    raise ValueError('Unsupported equation operation: ' + key)


def math_children(value):
    node = math_node(value)
    if node.tag == qn('m:oMath'):
        raise ValueError('Raw OMML must be a complete equation block, not a nested expression')
    return list(node) if isinstance(value, dict) and 'seq' in value else [node]


def add_equation(doc, value):
    expr = math_node(value)
    for run in expr.iter(qn('m:r')):
        math_properties = run.find(qn('m:rPr'))
        if math_properties is None:
            math_properties = element('rPr')
            run.insert(0, math_properties)
        if math_properties.find(qn('m:nor')) is None:
            math_properties.append(element('nor'))
        properties = run.find(qn('w:rPr'))
        if properties is None:
            properties = OxmlElement('w:rPr')
            run.insert(1, properties)
        set_run_fonts(properties)
    style = 'PLR Equation' if 'PLR Equation' in doc.styles else None
    if expr.tag == qn('m:oMath'):
        doc.add_paragraph(style=style)._p.append(element('oMathPara', [expr]))
        return
    math = element('oMath')
    if isinstance(value, dict) and 'seq' in value:
        for child in list(expr):
            math.append(child)
    else:
        math.append(expr)
    doc.add_paragraph(style=style)._p.append(element('oMathPara', [math]))


def audit_document(path):
    from zipfile import ZipFile
    with ZipFile(path) as archive:
        xml = etree.fromstring(archive.read('word/document.xml'))
        media = [n for n in archive.namelist() if n.startswith('word/media/')]
    invalid = xml.findall('.//m:oMath/m:e', NS) + xml.findall('.//m:e/m:e', NS) + xml.findall('.//m:num/m:e', NS) + xml.findall('.//m:den/m:e', NS)
    section_drawings = {str(i): 0 for i in range(1, 9)}
    section_equations = {str(i): 0 for i in range(1, 9)}
    current = None
    for p in xml.findall('./w:body/w:p', NS):
        text = ''.join(t.text or '' for t in p.findall('.//w:t', NS))
        if text in HEADINGS:
            current = str(HEADINGS.index(text) + 1)
        if current:
            section_drawings[current] += len(p.findall('.//w:drawing', NS))
    # Include formulas in tables and content controls, in body order.
    current = None
    for p in xml.find('./w:body', NS).iter(qn('w:p')):
        text = ''.join(t.text or '' for t in p.findall('.//w:t', NS))
        if text in HEADINGS:
            current = str(HEADINGS.index(text) + 1)
        if current:
            section_equations[current] += len(p.findall('.//m:oMathPara/m:oMath', NS))
    document = Document(path)
    resolutions = []
    for shape in document.inline_shapes:
        blips = shape._inline.xpath('.//a:blip')
        if not blips or not blips[0].get(qn('r:embed')):
            continue
        blob = document.part.related_parts[blips[0].get(qn('r:embed'))].blob
        image = Image.from_blob(blob)
        resolutions.append({'sha256': hashlib.sha256(blob).hexdigest(),
                            'metadata_dpi': min(image.horz_dpi, image.vert_dpi),
                            'effective_dpi': min(image.px_width / (shape.width / 914400),
                                                 image.px_height / (shape.height / 914400))})
    identifier = document.core_properties.identifier
    parts = identifier.split('|')
    manifest = ({'version': 1, 'stage': 'draft' if parts[1] == 'd' else 'assembled',
                 'input_sha256': parts[2], 'bound_sha256': parts[3], 'artifact_sha256': parts[4]}
                if len(parts) == 5 and parts[0] == 'plr1' and parts[1] in {'a', 'd'} else None)
    stored = document.settings.element.find('./w:docVars/w:docVar[@w:name="physics-lab-report-manifest"]', NS)
    if stored is not None:
        manifest = json.loads(stored.get(qn('w:val')))
    plain_letters = sum(bool(re.search(r'[A-Za-z\u00c0-\u024f\u0370-\u03ff]', t.text or ''))
                        for t in xml.findall('.//w:t', NS))
    punctuation_issues = 0
    for t in xml.findall('.//w:t', NS) + xml.findall('.//m:t', NS):
        if any(unicodedata.category(c).startswith('P') for c in t.text or ''):
            fonts = t.getparent().find('w:rPr/w:rFonts', NS)
            if fonts is None or any(fonts.get(qn('w:' + key)) != '宋体' for key in ('ascii', 'hAnsi', 'cs', 'eastAsia')):
                punctuation_issues += 1
    return {'unformatted_letters': plain_letters, 'punctuation_font_issues': punctuation_issues,
            'assembly_manifest': manifest, 'equations': len(xml.findall('.//m:oMath', NS)), 'invalid_math_containers': len(invalid),
            'section_equation_placements': section_equations,
            'section_image_placements': section_drawings,
            'image_placements': len(xml.findall('.//w:drawing', NS)),
            'embedded_media': len(media), 'image_resolutions': resolutions}


def payload_digest(payload):
    return hashlib.sha256(json.dumps(payload, sort_keys=True, ensure_ascii=False, allow_nan=False).encode('utf-8')).hexdigest()


def file_digest(file):
    digest = hashlib.sha256()
    with Path(file).open('rb') as f:
        for chunk in iter(lambda: f.read(1024*1024), b''):
            digest.update(chunk)
    return digest.hexdigest()


def template_path(payload, template=None):
    declared = payload.get('template_file')
    actual = Path(template or declared or TEMPLATE).resolve()
    if declared and Path(declared).resolve() != actual:
        raise ValueError('Template argument conflicts with template_file in report JSON')
    return actual


def assembly_manifest(payload, bound, draft=False, template=None):
    artifacts = list(payload.get('data_photos', [])) + list(payload.get('visual_sources', [])) + list(payload.get('data_files', []))
    artifacts += list(payload.get('result_files', {}).values())
    actual_template = template_path(payload, template)
    artifacts.append(str(actual_template))
    artifacts += [b['figure']['path'] for blocks in payload['sections'].values() for b in blocks if 'figure' in b]
    return {'version': 2, 'stage': 'draft' if draft else 'assembled',
            'template_file': str(actual_template),
            'input_sha256': payload_digest(payload), 'bound_sha256': payload_digest(bound),
            'artifact_sha256': payload_digest({str(Path(p).resolve()): file_digest(p) for p in artifacts})}


def bind_results(payload):
    """Resolve numerical placeholders from actual calculation JSON without mutating input."""
    payload = copy.deepcopy(payload)
    datasets = {key: json.loads(Path(path).read_text(encoding='utf-8-sig'))
                for key, path in payload.get('result_files', {}).items()}
    pattern = re.compile(r'\{\{result:([A-Za-z0-9_.-]+)(?::(\.[0-9]{1,2}[fge]))?\}\}')
    def replacement(match):
        parts = match[1].split('.')
        try:
            value = datasets[parts[0]]
            for key in parts[1:]:
                value = value[key]
        except (KeyError, TypeError) as exc:
            raise ValueError('Unknown calculation result ' + match[1]) from exc
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise ValueError('Result placeholder requires a finite number: ' + match[1])
        return format(value, match[2] or '.6g')
    def visit(value):
        if isinstance(value, str):
            result = pattern.sub(replacement, value)
            if '{{result:' in result:
                raise ValueError('Invalid or unresolved result placeholder')
            return result
        if isinstance(value, list):
            return [visit(x) for x in value]
        if isinstance(value, dict):
            return {key: visit(x) for key, x in value.items()}
        return value
    payload['sections'] = visit(payload['sections'])
    return payload


def content_issues(payload):
    if not isinstance(payload.get('explanation_notes', []), list) or any(not isinstance(n, str) for n in payload.get('explanation_notes', [])):
        raise ValueError('explanation_notes must be a list of strings for the separate delivery document')
    if 'questions_provided' in payload and type(payload['questions_provided']) is not bool:
        raise ValueError('questions_provided must be a boolean')
    sections = payload.get('sections')
    if not isinstance(sections, dict) or set(sections) != {str(i) for i in range(1, 9)}:
        raise ValueError('sections must contain keys 1 through 8')
    meaningful = {}
    for key, blocks in sections.items():
        if not isinstance(blocks, list):
            raise ValueError('Section blocks must be lists')
        count = 0
        for block in blocks:
            if not isinstance(block, dict) or len(block) != 1:
                raise ValueError('Every block must contain exactly one content key')
            kind, value = next(iter(block.items()))
            if kind == 'equation':
                math_node(value)
                count += 1
            elif kind in {'text', 'emphasis'}:
                if not isinstance(value, str):
                    raise ValueError('Text blocks must contain strings')
                count += bool(value.strip() and value.strip() not in {'待补充', '【待补充】', 'TODO', 'TBD'})
            elif kind == 'figure':
                if key == '5':
                    raise ValueError('Use data_photos for section 5')
                if not isinstance(value, dict) or not value.get('path'):
                    raise ValueError('Figure requires a path')
                chinese_label(value.get('caption'), 'Figure caption')
                chinese_label(value.get('note'), 'Figure note')
                count += 1
            elif kind == 'heading':
                if (not isinstance(value, dict) or not isinstance(value.get('text'), str) or not value['text'].strip()
                    or type(value.get('level')) is not int or not 1 <= value['level'] <= 5):
                    raise ValueError('Heading requires non-empty text and an integer level from 1 to 5')
                # A heading alone does not constitute substantive section content.
            elif kind == 'table':
                if not isinstance(value, dict) or not value.get('headers') or not isinstance(value.get('rows'), list):
                    raise ValueError('Table requires headers and rows')
                if any(not isinstance(row, list) or len(row) != len(value['headers']) for row in value['rows']):
                    raise ValueError('Table rows must match header width')
                chinese_label(value.get('caption'), 'Table caption')
                chinese_label(value.get('note'), 'Table note')
                count += bool(value['rows'] and any(str(cell).strip() for row in value['rows'] for cell in row))
            else:
                raise ValueError('Unsupported content block: ' + kind)
        meaningful[key] = count
    missing = list(payload.get('missing', []))
    if not payload.get('data_photos'):
        missing.append('原始实验数据照片')
    for key in ('6', '7', '8'):
        if key == '8' and payload.get('questions_provided') is False:
            continue
        if not meaningful[key]:
            missing.append(HEADINGS[int(key)-1] + '有效内容')
    if not any('equation' in b for b in sections['6']):
        missing.append('第六部分计算公式')
    if not any('figure' in b for b in sections['6']):
        missing.append('第六部分数据图')
    if not any(any(key in b and b[key].strip() for key in ('text', 'emphasis')) or ('table' in b and any(str(c).strip() for row in b['table']['rows'] for c in row)) for b in sections['6']):
        missing.append('第六部分计算过程')
    return list(dict.fromkeys(missing))


def build(payload, output, template=None, draft=False):
    output, template = Path(output).resolve(), template_path(payload, template)
    protected = [template] + list(payload.get('data_photos', [])) + list(payload.get('visual_sources', []))
    protected += list(payload.get('data_files', [])) + list(payload.get('result_files', {}).values())
    protected += [b['figure']['path'] for blocks in payload.get('sections', {}).values() for b in blocks if 'figure' in b]
    if any(output == Path(p).resolve() or (output.exists() and Path(p).exists() and output.samefile(p)) for p in protected):
        raise ValueError('Output must not overwrite an input file or template')
    if not isinstance(payload.get('experiment'), str) or not payload['experiment'].strip():
        raise ValueError('Missing experiment name')
    input_digest = payload_digest(payload)
    original_payload = payload
    payload = bind_results(payload)
    sections = payload['sections']
    photos = [Path(p).resolve() for p in payload.get('data_photos', [])]
    missing = content_issues(payload)
    if missing and not draft:
        raise ValueError('Incomplete report; use --draft or supply: ' + '、'.join(missing))
    images = list(photos)
    for key, blocks in sections.items():
        if not isinstance(blocks, list):
            raise ValueError('Section blocks must be lists')
        for block in blocks:
            if not isinstance(block, dict) or len(block) != 1:
                raise ValueError('Every block must contain exactly one content key')
            kind, value = next(iter(block.items()))
            if kind == 'equation':
                math_node(value)
            elif kind == 'figure':
                images.append(Path(value['path']).resolve())
                if key == '5':
                    raise ValueError('Use data_photos for section 5; figures belong elsewhere')
            elif kind == 'table':
                cols = len(value['headers'])
                if not cols or any(len(row) != cols for row in value['rows']):
                    raise ValueError('Table rows must match the header width')
            elif kind not in {'text', 'emphasis', 'heading'}:
                raise ValueError('Unsupported content block: ' + kind)
    for image in images:
        if not image.is_file():
            raise ValueError('Missing image: ' + str(image))
        if output == image:
            raise ValueError('Output must not overwrite an input image')
    doc = Document(template)
    # This asset is a structural reference; replace example drawings and text.
    configure_styles(doc)
    for child in list(doc._element.body):
        if child.tag != qn('w:sectPr'):
            doc._element.body.remove(child)
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.core_properties.title = payload['experiment'] + '实验报告'
    manifest = assembly_manifest(original_payload, payload, draft, template=template)
    set_assembly_manifest(doc, manifest)
    doc.add_paragraph('物理实验报告', style='Title')
    doc.add_paragraph('实验名称：' + payload['experiment'], style='PLR Info')
    info = doc.add_paragraph(style='PLR Info')
    for n, label in enumerate(('班级', '姓名', '学号', '座位号')):
        info.add_run(('  ' if n else '') + label + '：' + ' ' * 6)
    section = doc.sections[0]
    available = (section.page_width - section.left_margin - section.right_margin) / 914400
    picture_width = Inches(min(6.0, available))
    height_limit = min(7.5, (section.page_height - section.top_margin - section.bottom_margin) / 914400 - 1.0)

    def picture(path, caption, max_width=6.0, source=None, note=None, width_inches=None, max_height_inches=5.5):
        for number in (width_inches, max_height_inches):
            if number is not None and (isinstance(number, bool) or not isinstance(number, (int, float)) or not math.isfinite(number) or number <= 0):
                raise ValueError('Figure dimensions must be positive finite inches')
        p = doc.add_paragraph(style='PLR Figure')
        image = Image.from_file(str(path))
        width = min(picture_width, Inches(width_inches or max_width),
                    Inches(min(height_limit, max_height_inches) * image.px_width / image.px_height))
        if source == 'imagegen' and (min(image.horz_dpi, image.vert_dpi) < 300 or
                                     image.px_width / (width / 914400) < 300):
            raise ValueError('Imagegen figure requires metadata and effective resolution of at least 300 dpi')
        p.add_run().add_picture(str(path), width=width, height=int(width * image.px_height / image.px_width))
        p.paragraph_format.keep_with_next = True
        p = doc.add_paragraph(caption, style='PLR Caption')
        p.paragraph_format.keep_with_next = bool(note)
        if note:
            doc.add_paragraph(note, style='PLR Note')

    figure_number = table_number = 0
    for index, heading in enumerate(HEADINGS, 1):
        p = doc.add_paragraph(heading, style='Heading 1')
        p.paragraph_format.keep_with_next = True
        if index == 5:
            for n, photo in enumerate(photos, 1):
                picture(photo, '原始实验数据照片 ' + str(n))
        for block in sections[str(index)]:
            kind, value = next(iter(block.items()))
            if kind in {'text', 'emphasis'}:
                for line in str(value).splitlines():
                    if line.strip():
                        doc.add_paragraph(line, style='PLR Emphasis' if kind == 'emphasis' else 'Normal')
            elif kind == 'heading':
                doc.add_paragraph(value['text'], style='Heading ' + str(value['level']))
            elif kind == 'equation':
                add_equation(doc, value)
            elif kind == 'figure':
                figure_number += 1
                picture(value['path'], value.get('caption') or '图' + str(figure_number) + ' 实验数据图',
                        max_width=5.2, source=value.get('source'), note=value.get('note'),
                        width_inches=value.get('width_inches'), max_height_inches=value.get('max_height_inches', 4.6))
            elif kind == 'table':
                table_number += 1
                caption = doc.add_paragraph(value.get('caption') or '表' + str(table_number) + ' 实验数据', style='PLR Caption')
                caption.paragraph_format.keep_with_next = True
                table = doc.add_table(rows=1, cols=len(value['headers']))
                for cell, text in zip(table.rows[0].cells, value['headers']):
                    cell.text = str(text)
                for row in value['rows']:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = str(text)
                three_line_table(table, has_note=bool(value.get('note')))
                if value.get('note'):
                    doc.add_paragraph(value['note'], style='PLR Note')
    output.parent.mkdir(parents=True, exist_ok=True)
    format_document_text(doc)
    doc.save(output)
    return {'status': 'draft' if draft else 'assembled', 'input_sha256': input_digest, 'layout_verified': False,
            'missing': list(dict.fromkeys(missing)), 'delivery_notes': delivery_notes(payload) +
            (['草稿；待补充：' + '、'.join(dict.fromkeys(missing))] if draft else []),
            'data_photos': len(photos), **audit_document(output)}


def delivery_notes(payload):
    return (['未提供思考题'] if payload.get('questions_provided') is False else []) + list(payload.get('explanation_notes', []))


def set_assembly_manifest(doc, manifest):
    """Used by a generator when saving a newly assembled document, before review."""
    doc.core_properties.identifier = '|'.join(['plr1', 'd' if manifest['stage'] == 'draft' else 'a',
        manifest['input_sha256'], manifest['bound_sha256'], manifest['artifact_sha256']])
    # Document variables have room for the actual template path and version metadata.
    variables = doc.settings.element.find(qn('w:docVars'))
    if variables is None:
        variables = OxmlElement('w:docVars')
        doc.settings.element.insert_element_before(variables,
            'w:rsids', 'm:mathPr', 'w:attachedSchema', 'w:themeFontLang', 'w:clrSchemeMapping',
            'w:doNotIncludeSubdocsInStats', 'w:doNotAutoCompressPictures', 'w:forceUpgrade',
            'w:captions', 'w:readModeInkLockDown', 'w:smartTagType', 'sl:schemaLibrary',
            'w:shapeDefaults', 'w:doNotEmbedSmartTags', 'w:decimalSymbol', 'w:listSeparator')
    for existing in list(variables):
        if existing.get(qn('w:name')) == 'physics-lab-report-manifest':
            variables.remove(existing)
    variable = OxmlElement('w:docVar')
    variable.set(qn('w:name'), 'physics-lab-report-manifest')
    variable.set(qn('w:val'), json.dumps(manifest, ensure_ascii=False))
    variables.append(variable)


def load_payload(file):
    file = Path(file).resolve()
    payload = json.loads(file.read_text(encoding='utf-8-sig'))
    for key, value in payload.get('result_files', {}).items():
        payload['result_files'][key] = str(file.parent / value)
    for key in ('data_photos', 'visual_sources', 'data_files'):
        if key in payload:
            payload[key] = [str(file.parent / value) for value in payload[key]]
    if payload.get('template_file'):
        payload['template_file'] = str(file.parent / payload['template_file'])
    for blocks in payload['sections'].values():
        for block in blocks:
            if 'figure' in block:
                block['figure']['path'] = str(file.parent / block['figure']['path'])
    return payload


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('input', type=Path, help='Report JSON; relative image paths resolve beside JSON')
    parser.add_argument('output', type=Path)
    parser.add_argument('--template', type=Path)
    parser.add_argument('--draft', action='store_true')
    args = parser.parse_args()
    if args.input.resolve() == args.output.resolve():
        parser.error('Output must not overwrite input JSON')
    payload = load_payload(args.input)
    print(json.dumps(build(payload, args.output, args.template, args.draft), ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
