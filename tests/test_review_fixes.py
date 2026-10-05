"""Regression tests for protected inputs, actual templates and formula presence."""
import copy
import sys
import tempfile
import unittest
from pathlib import Path
from page_fixtures import synthetic_pages

from docx import Document
from docx.shared import Inches
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'physics-lab-report' / 'scripts'))
import build_report as b
import finalize_report as f


class ReviewFixTests(unittest.TestCase):
    def fixture(self, root):
        image = root / 'photo.png'
        Image.new('RGB', (100, 100)).save(image)
        payload = {'experiment': '回归测试', 'questions_provided': False,
                   'data_photos': [str(image)],
                   'sections': {str(i): [] for i in range(1, 9)}}
        payload['sections']['6'] = [{'text': '计算过程'}, {'equation': {'frac': ['U', 'I']}},
                                    {'figure': {'path': str(image)}}]
        payload['sections']['7'] = [{'text': '结果分析'}]
        return payload, image

    def passed_fixture(self, payload, output, root, image):
        # Synthetic records exercise the checker; no real review is claimed.
        quality = f.initialize_quality(payload, output, root, pages=synthetic_pages(root), code=[str(Path(__file__))])
        for record in quality['vision_reviews']:
            record.update(status='passed', checks=['synthetic fixture'], evidence='synthetic fixture')
        for record in quality['artifact_checks']:
            record.update(status='passed', evidence='synthetic fixture')
        for key in ('calculation_review', 'layout_review'):
            quality[key].update(status='passed', evidence='synthetic fixture')
        quality['execution_records'][0].update(exit_code=0, command='synthetic fixture')
        return quality

    def test_report_page_range_blocks_short_and_long_reports(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            payload, image = self.fixture(root)
            output = root / 'report.docx'
            b.build(payload, output)
            for count in (9, 10, 20, 21):
                with self.subTest(pages=count):
                    quality = self.passed_fixture(payload, output, root, image)
                    pages = []
                    for index in range(count):
                        page = root / ('page-' + str(index) + '.png')
                        Image.new('RGB', (32, 32), (index, 0, 0)).save(page)
                        pages.append(str(page))
                        quality['vision_reviews'].append({'path': str(page), 'sha256': f.digest_file(page),
                            'reviewer': 'current-model', 'status': 'passed', 'checks': ['synthetic fixture'],
                            'evidence': 'synthetic page-count contract only'})
                    quality['layout_review'].update(pages=pages, page_count=count)
                    result = f.evaluate(payload, output, quality, root)
                    self.assertEqual(result['status'], 'submission-ready' if 10 <= count <= 20 else 'review-pending')
                    self.assertEqual(any('10–20' in issue for issue in result['issues']), count in (9, 21))

    def test_output_cannot_overwrite_any_declared_input(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            payload, _ = self.fixture(root)
            for field in ('data_files', 'result_files', 'visual_sources', 'template_file'):
                with self.subTest(field=field):
                    source = root / (field + '.docx')
                    if field == 'template_file':
                        Document().save(source)
                    else:
                        source.write_text('{}', encoding='utf-8')
                    original = source.read_bytes()
                    case = copy.deepcopy(payload)
                    case[field] = ({'fit': str(source)} if field == 'result_files' else
                                   str(source) if field == 'template_file' else [str(source)])
                    with self.assertRaises(ValueError):
                        b.build(case, source)
                    self.assertEqual(source.read_bytes(), original)

    def test_actual_custom_template_change_blocks_finalization(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            payload, image = self.fixture(root)
            template = root / 'custom.docx'
            Document().save(template)
            output = root / 'report.docx'
            b.build(payload, output, template=template)
            quality = self.passed_fixture(payload, output, root, image)
            self.assertEqual(f.evaluate(payload, output, quality, root)['status'], 'submission-ready')
            doc = Document(template)
            doc.sections[0].left_margin = Inches(1.7)
            doc.save(template)
            # Even refreshing a file review must not bless a report from an old template.
            for record in quality['artifact_checks']:
                if Path(record['path']) == template:
                    record['sha256'] = f.digest_file(template)
            result = f.evaluate(payload, output, quality, root)
            self.assertEqual(result['status'], 'review-pending')
            self.assertTrue(any('版本' in issue for issue in result['issues']))

    def test_json_template_is_used_and_conflicting_argument_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            payload, _ = self.fixture(root)
            template = root / 'custom.docx'
            doc = Document()
            doc.sections[0].left_margin = Inches(1.7)
            doc.save(template)
            payload['template_file'] = str(template)
            output = root / 'report.docx'
            b.build(payload, output)
            self.assertEqual(Document(output).sections[0].left_margin, Inches(1.7))
            with self.assertRaises(ValueError):
                b.build(payload, root / 'conflicting.docx', template=b.TEMPLATE)

    def test_missing_or_moved_formula_blocks_finalization(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            payload, image = self.fixture(root)
            output = root / 'report.docx'
            for action in ('delete', 'move'):
                with self.subTest(action=action):
                    b.build(payload, output)
                    quality = self.passed_fixture(payload, output, root, image)
                    self.assertEqual(f.evaluate(payload, output, quality, root)['status'], 'submission-ready')
                    doc = Document(output)
                    equation = next(p._p for p in doc.paragraphs if p._p.findall('.//m:oMath', b.NS))
                    equation.getparent().remove(equation)
                    if action == 'move':
                        doc._element.body.insert(len(doc._element.body) - 1, equation)
                    doc.save(output)
                    quality['report_sha256'] = f.digest_file(output)
                    result = f.evaluate(payload, output, quality, root)
                    self.assertEqual(result['status'], 'review-pending')
                    self.assertTrue(any('公式' in issue for issue in result['issues']))


if __name__ == '__main__':
    unittest.main()
