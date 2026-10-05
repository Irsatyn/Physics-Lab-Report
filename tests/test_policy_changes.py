"""Regression cases for report completion and editable complex equations."""
import sys
import tempfile
import unittest
from pathlib import Path
from page_fixtures import synthetic_pages

from docx import Document
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'physics-lab-report' / 'scripts'))
import build_report as b
import finalize_report as f


class PolicyTests(unittest.TestCase):
    def payload(self, image):
        return {'experiment': '测试', 'data_photos': [str(image)],
                'sections': {str(i): [{'text': '内容'}] for i in range(1, 9)}}

    def test_no_questions_completes_without_notice_in_document(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / 'photo.png'
            Image.new('RGB', (600, 300)).save(image)
            payload = self.payload(image)
            payload['questions_provided'] = False
            payload['sections']['8'] = []
            payload['sections']['6'] += [{'equation': 'R=1'}, {'figure': {'path': str(image)}}]
            audit = b.build(payload, root / 'report.docx')
            text = '\n'.join(p.text for p in Document(root / 'report.docx').paragraphs)
            self.assertEqual(audit['status'], 'assembled')
            self.assertIn('未提供思考题', audit['delivery_notes'])
            self.assertNotIn('未提供思考题', text)
            self.assertNotIn('待补充', text)
            self.assertIn('八、课后思考题及实验拓展', text)
            payload['questions_provided'] = True
            with self.assertRaises(ValueError):
                b.build(payload, root / 'incomplete.docx')

    def test_complex_omml_is_native_and_invalid_xml_is_rejected(self):
        xml = '<m:oMath xmlns:m="' + b.NS['m'] + '"><m:m><m:mr><m:e><m:r><m:t>1</m:t></m:r></m:e></m:mr></m:m></m:oMath>'
        with tempfile.TemporaryDirectory() as tmp:
            doc = Document()
            b.add_equation(doc, {'omml': xml})
            path = Path(tmp) / 'matrix.docx'
            doc.save(path)
            self.assertEqual(b.audit_document(path)['equations'], 1)
            self.assertEqual(b.audit_document(path)['invalid_math_containers'], 0)
        for value in ['<x/>', '<m:oMath xmlns:m="' + b.NS['m'] + '"/>',
                      '<!DOCTYPE x [<!ENTITY x "bad">]>' + xml]:
            with self.subTest(value=value), self.assertRaises(ValueError):
                b.add_equation(Document(), {'omml': value})
        with self.assertRaises(ValueError):
            b.add_equation(Document(), {'seq': [{'omml': xml}]})

    def test_custom_template_version_is_bound_to_generated_report(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            template = root / 'custom.docx'
            Document().save(template)
            image = root / 'image.png'
            Image.new('RGB', (100, 100)).save(image)
            payload = self.payload(image)
            payload['template_file'] = str(template)
            payload['sections']['6'] += [{'equation': 'R=1'}, {'figure': {'path': str(image)}}]
            bound = b.bind_results(payload)
            original = b.assembly_manifest(payload, bound)
            doc = Document(template)
            doc.add_paragraph('custom generation fixture')
            b.set_assembly_manifest(doc, original)
            doc.save(root / 'custom-report.docx')
            self.assertEqual(b.audit_document(root / 'custom-report.docx')['assembly_manifest'], original)
            doc = Document(template)
            doc.add_paragraph('changed template')
            doc.save(template)
            self.assertNotEqual(b.assembly_manifest(payload, bound)['artifact_sha256'], original['artifact_sha256'])

    def test_imagegen_requires_real_300_dpi_at_inserted_size(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / 'plot.png'
            payload = self.payload(image)
            payload['sections']['6'] += [{'equation': 'R=1'},
                {'figure': {'path': str(image), 'source': 'imagegen'}}]
            for pixels, dpi in [(1024, 300), (1800, 72)]:
                Image.new('RGB', (pixels, pixels // 2)).save(image, dpi=(dpi, dpi))
                with self.subTest(pixels=pixels, dpi=dpi), self.assertRaises(ValueError):
                    b.build(payload, root / 'bad.docx')
            Image.new('RGB', (1800, 900)).save(image, dpi=(300, 300))
            b.build(payload, root / 'good.docx')
            audit = b.audit_document(root / 'good.docx')
            self.assertTrue(all(r['effective_dpi'] >= 300 for r in audit['image_resolutions']))

    def test_finalize_allows_no_questions_but_rejects_enlarged_imagegen(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            image = root / 'plot.png'
            Image.new('RGB', (1800, 900)).save(image, dpi=(300, 300))
            payload = self.payload(image)
            payload['questions_provided'] = False
            payload['sections']['8'] = []
            payload['sections']['6'] += [{'equation': 'R=1'},
                {'figure': {'path': str(image), 'source': 'imagegen'}}]
            output = root / 'report.docx'
            b.build(payload, output)
            quality = f.initialize_quality(payload, output, root, pages=synthetic_pages(root), code=[str(Path(__file__))])
            # Synthetic review evidence is a checker fixture, never a real report review.
            for record in quality['vision_reviews']:
                record.update(status='passed', checks=['unit-test fixture'], evidence='unit-test fixture')
            for key in ('calculation_review', 'layout_review'):
                quality[key].update(status='passed', evidence='unit-test fixture')
            quality['execution_records'][0].update(exit_code=0, command='unit-test fixture')
            result = f.evaluate(payload, output, quality, root)
            self.assertEqual(result['status'], 'submission-ready')
            self.assertEqual(result['delivery_notes'], ['未提供思考题'])
            doc = Document(output)
            shape = doc.inline_shapes[-1]
            shape.width = 7 * 914400
            shape.height = int(3.5 * 914400)
            doc.save(output)
            quality['report_sha256'] = f.digest_file(output)
            result = f.evaluate(payload, output, quality, root)
            self.assertEqual(result['status'], 'review-pending')
            self.assertTrue(any('300 dpi' in message for message in result['issues']))


if __name__ == '__main__':
    unittest.main()
