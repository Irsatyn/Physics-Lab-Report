"""User-facing Word typography and scientific table layout contracts."""
import sys
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Pt
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'physics-lab-report/scripts'))
import build_report as b


class FormattingTests(unittest.TestCase):
    def fixture(self, root):
        photo = root / 'photo.png'
        Image.new('RGB', (600, 300)).save(photo)
        payload = {'experiment': '排版测试', 'questions_provided': False, 'data_photos': [str(photo)],
                   'sections': {str(i): [] for i in range(1, 9)}}
        payload['sections']['6'] = [{'text': '正文包含中文和 U / V。'}, {'equation': {'frac': ['U', 'I']}},
            {'figure': {'path': str(photo), 'caption': '图1 测量值与拟合线', 'note': '注：误差棒表示标准不确定度。'}},
            {'table': {'caption': '表1 测量数据', 'note': '注：此表为合成测试数据。',
                       'headers': ['电流 / mA', '电压 / V'], 'rows': [[1, 3], [2, 5]]}}]
        payload['sections']['7'] = [{'text': '结果分析。'}]
        return payload

    def test_typography_alignment_and_four_plain_information_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b.build(self.fixture(root), root / 'report.docx')
            doc = Document(root / 'report.docx')
            normal = doc.styles['Normal']
            fonts = normal.element.rPr.find(qn('w:rFonts'))
            self.assertEqual(fonts.get(qn('w:eastAsia')), '宋体')
            self.assertEqual(fonts.get(qn('w:ascii')), 'Times New Roman')
            self.assertEqual(normal.font.size, Pt(12))
            body = next(p for p in doc.paragraphs if p.text.startswith('正文包含'))
            self.assertEqual(body.style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.JUSTIFY)
            self.assertEqual(body.style.element.pPr.find(qn('w:ind')).get(qn('w:firstLineChars')), '200')
            self.assertEqual(doc.paragraphs[0].style.font.size, Pt(22))
            heading = next(p for p in doc.paragraphs if p.text == b.HEADINGS[0])
            self.assertEqual(heading.style.font.size, Pt(15))
            self.assertEqual(heading.style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.LEFT)
            info = [p for p in doc.paragraphs if '班级：' in p.text]
            self.assertEqual(len(info), 1)
            self.assertEqual(info[0].text.replace(' ', ''), '班级：姓名：学号：座位号：')
            self.assertFalse(any('日期：' in p.text or '签章：' in p.text for p in doc.paragraphs))
            for p in info:
                self.assertEqual(p.style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
                self.assertFalse(any(run.underline for run in p.runs))

    def test_captions_notes_and_three_line_tables(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            b.build(self.fixture(root), root / 'report.docx')
            doc = Document(root / 'report.docx')
            title = next(p for p in doc.paragraphs if p.text == '图1 测量值与拟合线')
            self.assertEqual(title.style.font.size, Pt(10.5))
            self.assertTrue(title.style.font.bold)
            self.assertEqual(title.style.paragraph_format.alignment, WD_ALIGN_PARAGRAPH.CENTER)
            note = next(p for p in doc.paragraphs if p.text.startswith('注：误差棒'))
            self.assertEqual(note.style.font.size, Pt(9))
            self.assertTrue(note.style.font.bold)
            self.assertIs(title._p.getnext(), note._p)
            table = doc.tables[0]
            self.assertEqual(table.alignment, WD_TABLE_ALIGNMENT.CENTER)
            borders = table._tbl.tblPr.find(qn('w:tblBorders'))
            self.assertEqual(borders.find(qn('w:top')).get(qn('w:val')), 'single')
            self.assertEqual(borders.find(qn('w:bottom')).get(qn('w:val')), 'single')
            for edge in ('left', 'right', 'insideH', 'insideV'):
                self.assertEqual(borders.find(qn('w:' + edge)).get(qn('w:val')), 'nil')
            for cell in table.rows[0].cells:
                bottom = cell._tc.tcPr.find('./w:tcBorders/w:bottom', b.NS)
                self.assertEqual(bottom.get(qn('w:val')), 'single')
            self.assertEqual(table._tbl.getnext().tag, qn('w:p'))
            self.assertIn('此表为合成', ''.join(table._tbl.getnext().itertext()))
            self.assertIsNotNone(table.rows[0]._tr.trPr.find(qn('w:tblHeader')))
            self.assertTrue(table.rows[-1].cells[0].paragraphs[0].paragraph_format.keep_with_next)

    def test_nested_headings_use_size_then_bold_and_italic(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            payload = self.fixture(root)
            payload['sections']['6'] = [{'heading': {'text': '子标题' + str(level), 'level': level}}
                                      for level in (2, 3, 4, 5)] + payload['sections']['6']
            b.build(payload, root / 'report.docx')
            doc = Document(root / 'report.docx')
            for level, size, bold, italic in [(2, 14, True, False), (3, 12, True, False),
                                               (4, 12, False, True), (5, 12, True, True)]:
                p = next(p for p in doc.paragraphs if p.text == '子标题' + str(level))
                self.assertEqual(p.style.font.size, Pt(size))
                self.assertEqual(p.style.font.bold, bold)
                self.assertEqual(p.style.font.italic, italic)
                self.assertTrue(p.style.paragraph_format.keep_with_next)

    def test_explicit_english_only_caption_is_rejected(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            payload = self.fixture(root)
            for kind in ('figure', 'table'):
                with self.subTest(kind=kind):
                    value = next(block[kind] for block in payload['sections']['6'] if kind in block)
                    original = value['caption']
                    value['caption'] = 'Figure 1 Measurements'
                    with self.assertRaises(ValueError):
                        b.build(payload, root / 'invalid.docx')
                    value['caption'] = original


if __name__ == '__main__':
    unittest.main()
